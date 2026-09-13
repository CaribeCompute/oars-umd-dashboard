"""Read the supplied workbook without editing it. Run with Python + openpyxl."""
from pathlib import Path
import json, re, hashlib
from collections import Counter
from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/OARS Mid-Atlantic Tool Database.xlsx'
wb = load_workbook(SOURCE, data_only=True)
# Explicit mappings preserve the different worksheet schemas; no blanket fill-down.
common = dict(name='D', problem='B', stage='C', county='E', agency='F', website='G', land='H', eligibility='I', requirements='J', practice='K', practiceWebsite='L', strategies='M')
specs = {
 'Federal Program Overview': (8,'Federal',dict(name='D',problem='B',stage='C',agency='E',website='F',land='G',eligibility='H',requirements='I',practice='J',practiceWebsite='K',strategies='L',contacts='M',notes='N')),
 'Federal Practices': (9,'Federal',dict(name='E',problem='B',stage='C',programName='D',code='F',practiceWebsite='G',land='H',strategies='I',eligibility='J',requirements='K',notes='L')),
 'MD': (8,'MD',dict(common,contacts='N',notes='O')),
 'DE': (6,'DE',dict(common,costShare='O',benefit='P',personnel='Q',contacts='R',limitations='S',quantitative='T',qualitative='U',code='V')),
 'NJ': (8,'NJ',dict(common,costShare='N',benefit='O',personnel='P',contacts='Q')),
 'VA': (8,'VA',dict(common,costShare='N',benefit='O',personnel='P',contacts='Q')),
 'Private Program Overview': (8,'Private',dict(name='E',scopeDetail='B',problem='C',stage='D',agency='F',website='G',land='H',eligibility='I',requirements='J',practice='K',practiceWebsite='L',strategies='M',contacts='N',notes='O')),
 'SWI Shortlisted Programs & Prac': (4,'Federal',dict(scope='A',county='B',agency='C',programName='D',website='E',land='F',name='G',practiceWebsite='H',overviewUrl='I',description='J',stage='K',strategies='L',swiStrategies='M',nextStep='N',experts='O',problem='P',goals='Q',contacts='S',eligibility='T',requirements='U',costShare='V',benefit='W',timeline='X',duration='Y',limitations='Z',quantitative='AA',qualitative='AB',notes='AC')),
}
def clean(v):
 if v is None:return ''
 if isinstance(v,float) and v.is_integer():return str(int(v))
 return re.sub(r'file://\S+', '[local file reference omitted]', str(v).strip())
def get(sheet,row,col,use_link=False):
 cell=sheet[f'{col}{row}']
 # Only actual merged cells inherit their anchor; empty cells stay unknown.
 if cell.value is None:
  for merged in sheet.merged_cells.ranges:
   if cell.coordinate in merged:cell=sheet.cell(merged.min_row,merged.min_col);break
 if use_link and cell.hyperlink and cell.hyperlink.target and cell.hyperlink.target.startswith(("https://", "http://")):
  return clean(cell.hyperlink.target)
 return clean(cell.value)
records=[]
for sheet_name,(header,scope,cols) in specs.items():
 s=wb[sheet_name]
 for row in range(header+1,s.max_row+1):
  fields={key:get(s,row,col,key in ["website","practiceWebsite","overviewUrl"]) for key,col in cols.items()}
  if not fields['name']:continue
  # Numeric-only row counters and formatting rows never become records.
  if not any(fields.get(k) for k in ['website','practiceWebsite','problem','description','eligibility','agency','programName']):continue
  rawland=fields.get('land','').lower()
  farm=any(x in rawland for x in ['farm','agricultur','crop'])
  forest=any(x in rawland for x in ['forest','wood'])
  land='both' if farm and forest else 'farm' if farm else 'forest' if forest else 'unspecified'
  actual_scope=fields.pop('scope',scope)
  if actual_scope in ['US','U.S.','U.S']:actual_scope='Federal'
  kind='Practice' if sheet_name in ['Federal Practices','SWI Shortlisted Programs & Prac'] else 'Program'
  records.append(dict(id=f'{re.sub("[^a-z0-9]+","-",sheet_name.lower()).strip("-")}-{row}',name=fields.pop('name'),type=kind,land=land,landDescription=fields.pop('land',''),scope=actual_scope,shortlisted=sheet_name.startswith('SWI Shortlisted'),sourceSheet=sheet_name,sourceRow=row,**fields))
# The explicit exclusion sheet does not become an eligible program record.
s=wb['Programs- Not Applicable']
excluded=[{'name':clean(s[f'C{r}'].value),'reason':clean(s[f'F{r}'].value),'sourceRow':r} for r in range(5,s.max_row+1) if clean(s[f'C{r}'].value)]
excluded_names={e['name'].casefold() for e in excluded}
records=[r for r in records if r['name'].casefold() not in excluded_names]
output={'sourceFile':SOURCE.name,'sourceSha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'importedOn':'2026-09-13','records':records,'excluded':excluded,'countsBySheet':dict(Counter(r['sourceSheet'] for r in records))}
(ROOT/'dashboard/data/oars-programs.json').write_text(json.dumps(output,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'records':len(records),'counts':output['countsBySheet'],'excluded':len(excluded),'missingStage':sum(not r.get('stage') for r in records)},indent=2))
