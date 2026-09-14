"""Build the prototype catalog from the OARS workbook shortlist.

Run from dashboard/: python scripts/import-oars-workbook.py [workbook.xlsx]
Requires openpyxl. The generated TypeScript remains usable without Python.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


DASHBOARD = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = DASHBOARD.parents[1] / "docs" / "Copy of OARS Mid-Atlantic Tool Database.xlsx"
OUTPUT = DASHBOARD / "content" / "mid-atlantic-catalog.ts"


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def slug(value: str) -> str:
    return re.sub(r"(^-|-$)", "", re.sub(r"[^a-z0-9]+", "-", value.lower()))


def land_types(value: object) -> list[str]:
    value_text = text(value).lower()
    result: list[str] = []
    if "farm" in value_text:
        result.append("farm")
    if "forest" in value_text:
        result.append("forest")
    return result or ["farm", "forest"]


GOAL_KEYWORDS = {
    "agriculture": r"agricultur|crop|farm|forage|pasture|graz|production|soil health",
    "transition": r"salt.tolerant|salin|transition|convert|alternative vegetation",
    "habitat": r"habitat|wildlife|wetland|pollinator|riparian|buffer|biodiversity",
    "income": r"income|economic|payment|productivity|production|financial",
    "legacy": r"easement|legacy|preserv|long.term|protect.*property",
    "infrastructure": r"drain|flood|channel|ditch|dike|levee|road|structure for water control",
    "woodlot-management": r"forest|tree|shrub|wood|silvicultur|canopy",
    "invasive-species": r"invasive|brush|weed|phragmites|undesirable vegetation",
    "minimize-costs": r"cost.share|financial assistance|funding|payment|grant",
    "marsh-transition": r"marsh|wetland|tidal|hydrolog",
}

CONCERN_KEYWORDS = {
    "soil-health": r"soil|erosion|tillage|cover crop|nutrient|phosphorus|runoff|gypsum",
    "water-management": r"water|drain|irrigat|flood|wetland|channel|runoff|tidal",
    "habitat": r"habitat|wildlife|pollinator|hedgerow|buffer|tree|forest|wood",
    "vegetation-management": r"invasive|brush|woody|phragmites|vegetation|weed",
}


def matches(source: str, patterns: dict[str, str]) -> list[str]:
    return [key for key, pattern in patterns.items() if re.search(pattern, source, re.I)]


def build_catalog(workbook_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    shortlist = workbook["SWI Shortlisted Programs & Prac"]
    federal = workbook["Federal Program Overview"]

    eqip = next(
        row for row in federal.iter_rows(min_row=9, values_only=True)
        if text(row[3]) == "Environmental Quality Incentives Program (EQIP)"
    )
    program = {
        "id": "program-eqip",
        "name": text(eqip[3]),
        "agency": text(eqip[4]),
        "type": "Program",
        "landTypes": land_types(eqip[6]),
        "stageIds": [],
        "goalIds": [],
        "concernIds": [],
        "description": text(eqip[1]),
        "eligibility": text(eqip[7]) or None,
        "requirements": text(eqip[8]) or None,
        "costShare": None,
        "paymentBenefit": None,
        "timeline": None,
        "duration": None,
        "deadline": None,
        "limitations": text(eqip[13]) or None,
        "sourceUrl": text(eqip[5]) or None,
        "practiceUrl": None,
        "practiceOverviewUrl": None,
        "strategies": [],
        "nextStep": "Contact an NRCS service center to confirm current eligibility and ranking dates.",
        "contact": text(eqip[12]) or None,
        "scope": "United States",
        "county": None,
        "programId": None,
        "programName": None,
        "source": f"{workbook_path.name} / Federal Program Overview row 15",
        "sourceRow": 15,
        "mappingStatus": "oars-input-required",
        "status": "published",
    }

    practices: list[dict[str, object]] = []
    for row_number, row in enumerate(shortlist.iter_rows(min_row=5, max_row=250, values_only=True), 5):
        practice_name = text(row[6])
        if not practice_name:
            continue
        combined = " ".join(text(row[index]) for index in (5, 9, 11, 12, 15, 16, 19, 20, 25, 28))
        strategies = [item.strip() for item in re.split(r"[;\n]+", " ; ".join(filter(None, [text(row[11]), text(row[12])]))) if item.strip()]
        practices.append({
            "id": f"practice-{slug(practice_name)}",
            "name": practice_name,
            "agency": text(row[2]),
            "type": "Practice",
            "landTypes": land_types(row[5]),
            "stageIds": [],
            "goalIds": matches(combined, GOAL_KEYWORDS),
            "concernIds": matches(combined, CONCERN_KEYWORDS),
            "description": text(row[9]) or text(row[11]) or text(row[15]),
            "eligibility": text(row[19]) or None,
            "requirements": text(row[20]) or None,
            "costShare": text(row[21]) or None,
            "paymentBenefit": text(row[22]) or None,
            "timeline": text(row[23]) or None,
            "duration": text(row[24]) or None,
            "deadline": None,
            "limitations": text(row[25]) or text(row[28]) or None,
            "sourceUrl": text(row[4]) or None,
            "practiceUrl": text(row[7]) or None,
            "practiceOverviewUrl": text(row[8]) or None,
            "strategies": strategies,
            "nextStep": text(row[13]) or None,
            "contact": text(row[18]) or text(row[14]) or None,
            "scope": text(row[0]) or None,
            "county": text(row[1]) or None,
            "programId": "program-eqip",
            "programName": text(row[3]),
            "source": f"{workbook_path.name} / SWI Shortlisted Programs & Prac row {row_number}",
            "sourceRow": row_number,
            "mappingStatus": "provisional-keyword",
            "status": "published",
        })
    return [program, *practices]


def main() -> None:
    workbook_path = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    catalog = build_catalog(workbook_path)
    output = (
        "// Generated from the OARS workbook by scripts/import-oars-workbook.py.\n"
        "// Stage columns are blank in the source shortlist. Goal mappings are provisional keyword matches pending OARS approval.\n"
        f"const catalog = {json.dumps(catalog, ensure_ascii=False, indent=2)};\n\n"
        "export default catalog;\n"
    )
    OUTPUT.write_text(output, encoding="utf-8")
    print(f"Wrote {len(catalog)} resources ({len(catalog) - 1} practices) to {OUTPUT}")


if __name__ == "__main__":
    main()
